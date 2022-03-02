from pathlib import Path

import openpyxl


class ExcelDataLoader:
    wb: openpyxl.Workbook

    def __init__(self, path: Path):
        self.wb = openpyxl.load_workbook(str(path), data_only=True)
        self.all_tables: dict[str, list[str]] = {}
        self.table_ws_map: dict[str, str] = {}
        for ws in self.wb:
            for name in ws.tables:
                self.table_ws_map[name] = ws.title
                self.all_tables[name] = [
                    col.name for col in ws.tables[name].tableColumns
                ]

    def tables(self, config) -> list[str]:
        return list(self.all_tables)

    def headers(self, config) -> list[str]:
        return self.all_tables.get(config.table_name, [])

    def load_data(self, config) -> tuple[list[str], list[list]]:
        if config.table_name not in self.all_tables:
            return [], []
        data = [
            [cell.value for cell in row]
            for row in list(self.wb[self.table_ws_map[config.table_name]])[1:]
        ]
        return self.all_tables[config.table_name], data


def register():
    return (".xlsx", ".xlsm"), ExcelDataLoader
