from pathlib import Path

import openpyxl
from emsapp.config import Config
from emsapp.data import Importer


class ExcelImporter(Importer, ext=(".xlsx", ".xlsm")):
    wb: openpyxl.Workbook

    def __init__(self, path: Path):
        self.wb = openpyxl.load_workbook(str(path), data_only=True)
        self.all_tables: dict[str, list[str]] = {}
        self.table_ws_map: dict[str, str] = {}
        for ws in self.wb:
            for name in ws.tables:
                self.table_ws_map[name] = ws.title
                self.all_tables[name] = [col.name for col in ws.tables[name].tableColumns]

    def tables(self) -> list[str]:
        return list(self.all_tables)

    def headers(self) -> list[str]:
        table_name = Config().data.table_name
        return self.all_tables.get(table_name, [])

    def import_data(self) -> tuple[list[str], list[list]]:
        table_name = Config().data.table_name
        if table_name not in self.all_tables:
            return [], []
        data = [[cell.value for cell in row] for row in self.wb[self.table_ws_map[table_name]][1:]]
        return self.all_tables[table_name], data
